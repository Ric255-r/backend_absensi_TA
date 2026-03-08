import hashlib
import os
from collections import defaultdict
from datetime import date, datetime, time, timedelta
import json

from fastapi import HTTPException
from fastapi_jwt import JwtAuthorizationCredentials
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection

from app.models import (
  Akun,
  AuditLog,
  Departemen,
  HariLibur,
  JadwalKerja,
  Karyawan,
  KonfigurasiAplikasi,
  PengajuanAbsen,
)
from app.models.absensi import Absensi
from app.core.audit import audit_action, log_audit
from app.schemas.requests.admin import (
  AkunCreateRequest,
  AkunUpdateRequest,
  DepartemenCreateRequest,
  HariLiburCreateRequest,
  HariLiburUpdateRequest,
  JadwalUpdateRequest,
  KaryawanCreateRequest,
  KaryawanUpdateRequest,
  KonfigurasiUpdateRequest,
  UpdatePengajuanRequest,
  UpdateStatusAbsensiRequest,
)
from utils.fn_log import logger
from app.realtime.absensi_ws import admin_to_user_conn
from tortoise.transactions import in_transaction


def _resolve_analytics_range(
  start_date: str | None,
  end_date: str | None,
) -> tuple[date, date]:
  today = date.today()
  resolved_end = date.fromisoformat(end_date) if end_date else today
  resolved_start = (
    date.fromisoformat(start_date)
    if start_date
    else (resolved_end - timedelta(days=29))
  )
  if resolved_end < resolved_start:
    raise HTTPException(
      status_code=400, detail="Tanggal akhir tidak boleh < tanggal mulai"
    )
  return resolved_start, resolved_end


async def regis_karyawan(payload: KaryawanCreateRequest) -> dict:
  tanggal_rekrut = (
    date.fromisoformat(payload.tanggal_rekrut) if payload.tanggal_rekrut else None
  )

  await Karyawan.create(
    id_karyawan=payload.id_karyawan,
    nama_karyawan=payload.nama_karyawan,
    email_karyawan=payload.email_karyawan,
    nomor_hp=payload.nomor_hp,
    tanggal_rekrut=tanggal_rekrut,
    status=payload.status,
    departemen_id=payload.id_departemen,
    posisi=payload.posisi,
  )
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def regis_akun(payload: AkunCreateRequest) -> dict:
  passwd = hashlib.md5(str(payload.passwd).encode()).hexdigest()
  await Akun.create(
    username=payload.username,
    passwd=passwd,
    roles=payload.roles,
    karyawan_id=payload.id_karyawan,
    status=payload.status,
  )
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def regis_departemen(payload: DepartemenCreateRequest) -> dict:
  await Departemen.create(nama_departemen=payload.nama_departemen)
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def regis_hari_libur(payload: HariLiburCreateRequest) -> dict:
  await HariLibur.create(
    tanggal=date.fromisoformat(payload.tanggal),
    keterangan=payload.keterangan,
    tipe=payload.tipe,
  )
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def get_karyawan(id_karyawan: str | None = None):
  query = Karyawan.all()
  if id_karyawan:
    items = (
      await query.filter(id_karyawan=id_karyawan)
      .limit(1)
      .values(
        "id_karyawan",
        "nama_karyawan",
        "email_karyawan",
        "nomor_hp",
        "foto_profile",
        "tanggal_rekrut",
        "status",
        "posisi",
        "departemen_id",
        "departemen__nama_departemen",
      )
    )
    return items[0] if items else None
  return await query.values(
    "id_karyawan",
    "nama_karyawan",
    "email_karyawan",
    "nomor_hp",
    "foto_profile",
    "tanggal_rekrut",
    "status",
    "posisi",
    "departemen_id",
    "departemen__nama_departemen",
  )


async def get_akun(username: str | None = None):
  query = Akun.all()
  if username:
    items = (
      await query.filter(username=username)
      .limit(1)
      .values(
        "username",
        "roles",
        "status",
        "last_login",
        "device_id",
        "karyawan_id",
      )
    )
    item = items[0] if items else None
    return [item] if item else []
  return await query.values(
    "username",
    "roles",
    "status",
    "last_login",
    "device_id",
    "karyawan_id",
  )


async def get_departemen():
  return await Departemen.all().values("id_departemen", "nama_departemen")


async def get_jadwal():
  return await JadwalKerja.all().values(
    "id_jadwal", "hari_dalam_seminggu", "shift_mulai", "shift_selesai"
  )


async def get_konfigurasi():
  items = (
    await KonfigurasiAplikasi.all()
    .order_by("id_pengaturan")
    .limit(1)
    .values("id_pengaturan", "toleransi_terlambat", "maks_hari_cuti")
  )
  return items[0] if items else None


async def get_hari_libur():
  return (
    await HariLibur.all()
    .order_by("tanggal")
    .values("id_libur", "tanggal", "keterangan", "tipe")
  )


async def get_data_dashboard(tgl: str | None = None):
  target_date = date.fromisoformat(tgl) if tgl else date.today()
  day_start = datetime.combine(target_date, time.min)
  day_end = day_start + timedelta(days=1)

  total_karyawan = await Karyawan.filter(status="aktif").count()
  total_karyawan_rows = (
    await Karyawan.filter(status="aktif")
    .order_by("nama_karyawan")
    .values(
      "id_karyawan",
      "nama_karyawan",
      "email_karyawan",
      "nomor_hp",
      "foto_profile",
      "tanggal_rekrut",
      "status",
      "posisi",
      "departemen_id",
      "departemen__nama_departemen",
    )
  )
  total_karyawan_list = [
    {
      "id_karyawan": row["id_karyawan"],
      "nama_karyawan": row["nama_karyawan"],
      "email_karyawan": row["email_karyawan"],
      "nomor_hp": row["nomor_hp"],
      "foto_profile": row["foto_profile"],
      "tanggal_rekrut": row["tanggal_rekrut"],
      "status": row["status"],
      "posisi": row["posisi"],
      "departemen_id": row["departemen_id"],
      "nama_departemen": row["departemen__nama_departemen"],
    }
    for row in total_karyawan_rows
  ]

  absen_pending_count = await Absensi.filter(
    status_absen="pending",
    tanggal_absen__gte=day_start,
    tanggal_absen__lt=day_end,
  ).count()
  absen_pending_rows = (
    await Absensi.filter(
      status_absen="pending",
      tanggal_absen__gte=day_start,
      tanggal_absen__lt=day_end,
    )
    .order_by("-tanggal_absen")
    .values(
      "id_absensi",
      "id_karyawan",
      "tanggal_absen",
      "check_in",
      "check_out",
      "pengajuan",
      "status_absen",
      "karyawan__nama_karyawan",
      "karyawan__email_karyawan",
      "karyawan__nomor_hp",
      "karyawan__foto_profile",
      "karyawan__tanggal_rekrut",
      "karyawan__status",
      "karyawan__posisi",
      "karyawan__departemen_id",
    )
  )
  absen_pending_list = [
    {
      "id": row["id_absensi"],
      "id_karyawan": row["id_karyawan"],
      "id_absensi": row["id_absensi"],
      "tanggal_absen": row["tanggal_absen"],
      "check_in": row["check_in"],
      "check_out": row["check_out"],
      "pengajuan": row["pengajuan"],
      "status_absen": row["status_absen"],
      "nama_karyawan": row["karyawan__nama_karyawan"],
      "email_karyawan": row["karyawan__email_karyawan"],
      "nomor_hp": row["karyawan__nomor_hp"],
      "foto_profile": row["karyawan__foto_profile"],
      "tanggal_rekrut": row["karyawan__tanggal_rekrut"],
      "status": row["karyawan__status"],
      "posisi": row["karyawan__posisi"],
      "id_departemen": row["karyawan__departemen_id"],
    }
    for row in absen_pending_rows
  ]

  ga_hadir_count = await Absensi.filter(
    pengajuan__in=["cuti", "sakit", "izin"],
    tanggal_absen__gte=day_start,
    tanggal_absen__lt=day_end,
  ).count()
  ga_hadir_rows = (
    await Absensi.filter(
      pengajuan__in=["cuti", "sakit", "izin"],
      tanggal_absen__gte=day_start,
      tanggal_absen__lt=day_end,
    )
    .order_by("karyawan__nama_karyawan")
    .values(
      "id_absensi",
      "id_karyawan",
      "tanggal_absen",
      "pengajuan",
      "status_absen",
      "karyawan__nama_karyawan",
    )
  )
  ga_hadir_list = [
    {
      "id": row["id_absensi"],
      "id_karyawan": row["id_karyawan"],
      "nama": row["karyawan__nama_karyawan"],
      "tanggal_absen": row["tanggal_absen"],
      "pengajuan": row["pengajuan"],
      "status_absen": row["status_absen"],
    }
    for row in ga_hadir_rows
  ]

  logger.info("ADMIN MENGAKSES DASHBOARD")

  return {
    "total_karyawan": {"karyawan": total_karyawan},
    "total_karyawan_list": total_karyawan_list,
    "absen_pending": {"pending": absen_pending_count},
    "absen_pending_list": absen_pending_list,
    "data_ga_hadir": {"ga_hadir": ga_hadir_count},
    "ga_hadir_list": ga_hadir_list,
  }


async def get_analytics(start_date: str | None = None, end_date: str | None = None):
  start_d, end_d = _resolve_analytics_range(start_date, end_date)
  start_dt = datetime.combine(start_d, time.min)
  end_dt = datetime.combine(end_d + timedelta(days=1), time.min)
  days_count = (end_d - start_d).days + 1

  absensi_rows = await Absensi.filter(
    tanggal_absen__gte=start_dt,
    tanggal_absen__lt=end_dt,
  ).values(
    "id_absensi",
    "id_karyawan",
    "tanggal_absen",
    "status_absen",
    "pengajuan",
    "is_telat",
    "karyawan__nama_karyawan",
    "karyawan__departemen__nama_departemen",
  )
  pengajuan_rows = await PengajuanAbsen.filter(
    tanggal_mulai__lte=end_d,
    tanggal_akhir__gte=start_d,
  ).values(
    "id_pengajuan",
    "id_karyawan",
    "status",
    "tanggal_mulai",
    "tanggal_akhir",
    "karyawan__nama_karyawan",
  )
  audit_rows = (
    await AuditLog.filter(
      created_at__gte=start_dt,
      created_at__lt=end_dt,
    )
    .order_by("-created_at")
    .values(
      "id",
      "actor_username",
      "actor_id_karyawan",
      "actor_role",
      "action",
      "table_name",
      "record_id",
      "created_at",
      "metadata",
    )
  )

  trend_map: dict[str, dict] = {}
  for i in range(days_count):
    cur_day = start_d + timedelta(days=i)
    key = cur_day.isoformat()
    trend_map[key] = {
      "date": key,
      "hadir": 0,
      "tidak_hadir": 0,
      "pending": 0,
      "telat": 0,
    }

  dept_map: dict[str, dict] = {}
  late_by_employee: dict[str, dict] = {}
  absent_by_employee: dict[str, dict] = {}

  present_types = {"hadir"}
  absent_types = {"cuti", "sakit", "izin"}

  present_count = 0
  absent_count = 0
  pending_count = 0
  late_count = 0

  for row in absensi_rows:
    row_date = row["tanggal_absen"].date().isoformat()
    pengajuan = (row.get("pengajuan") or "hadir").lower()
    status_absen = (row.get("status_absen") or "").lower()
    is_telat = int(row.get("is_telat") or 0)
    name = row.get("karyawan__nama_karyawan") or "-"
    dep = row.get("karyawan__departemen__nama_departemen") or "Tanpa Departemen"

    if dep not in dept_map:
      dept_map[dep] = {
        "departemen": dep,
        "total_record": 0,
        "hadir": 0,
        "tidak_hadir": 0,
        "pending": 0,
        "telat": 0,
      }
    dept_map[dep]["total_record"] += 1

    if row_date in trend_map:
      if pengajuan in absent_types:
        trend_map[row_date]["tidak_hadir"] += 1
      else:
        trend_map[row_date]["hadir"] += 1
      if status_absen == "pending":
        trend_map[row_date]["pending"] += 1
      if is_telat == 1:
        trend_map[row_date]["telat"] += 1

    if pengajuan in absent_types:
      absent_count += 1
      dept_map[dep]["tidak_hadir"] += 1
      if name not in absent_by_employee:
        absent_by_employee[name] = {
          "id_karyawan": row["id_karyawan"],
          "nama_karyawan": name,
          "jumlah_tidak_hadir": 0,
        }
      absent_by_employee[name]["jumlah_tidak_hadir"] += 1
    elif pengajuan in present_types:
      present_count += 1
      dept_map[dep]["hadir"] += 1
    else:
      present_count += 1
      dept_map[dep]["hadir"] += 1

    if status_absen == "pending":
      pending_count += 1
      dept_map[dep]["pending"] += 1

    if is_telat == 1:
      late_count += 1
      dept_map[dep]["telat"] += 1
      if name not in late_by_employee:
        late_by_employee[name] = {
          "id_karyawan": row["id_karyawan"],
          "nama_karyawan": name,
          "jumlah_telat": 0,
        }
      late_by_employee[name]["jumlah_telat"] += 1

  departemen_breakdown = []

  for row in dept_map.values():
    hadir_base = row["hadir"] if row["hadir"] > 0 else 1
    row["late_rate_percent"] = round((row["telat"] / hadir_base) * 100, 2)
    departemen_breakdown.append(row)

  # Membuat fungsi bantuan untuk menggantikan 'lambda'
  def kriteria_sort_departemen(x):
    # Mengembalikan nilai untuk diurutkan
    return (x["late_rate_percent"], x["telat"])

  # Menggunakan reverse=True agar diurutkan dari nilai terbesar ke terkecil (Descending)
  departemen_breakdown.sort(key=kriteria_sort_departemen, reverse=True)

  # ---------------------------------------------------------
  # 2. MENCARI TOP 5 KARYAWAN TELAT & TIDAK HADIR
  # ---------------------------------------------------------
  def kriteria_sort_telat(x):
    # Tanda minus (-) berarti diurutkan menurun (angka terbesar dulu)
    # Nama karyawan tanpa minus berarti diurutkan menaik sesuai abjad A-Z
    return (-x["jumlah_telat"], x["nama_karyawan"])

  top_telat_semua = sorted(late_by_employee.values(), key=kriteria_sort_telat)
  top_telat = top_telat_semua[0:5]  # Mengambil 5 data pertama

  def kriteria_sort_absen(x):
    return (-x["jumlah_tidak_hadir"], x["nama_karyawan"])

  top_tidak_hadir_semua = sorted(absent_by_employee.values(), key=kriteria_sort_absen)
  top_tidak_hadir = top_tidak_hadir_semua[0:5]

  # ---------------------------------------------------------
  # 3. MEMFILTER APPROVAL LOGS
  # ---------------------------------------------------------
  approval_logs = []

  # Mengganti List Comprehension dengan for loop dan if standar
  for row in audit_rows:
    # Menggunakan list biasa [...] bukan set {...} agar lebih awam
    if row["table_name"] in ["absensi", "pengajuan_absen"]:
      if row["action"] in ["approve", "reject"]:
        approval_logs.append(row)

  # ---------------------------------------------------------
  # 4. MENGHITUNG TOTAL APPROVE DAN REJECT
  # ---------------------------------------------------------
  approved_actions = 0
  rejected_actions = 0

  # Mengganti fungsi sum() dan generator dengan cara hitung manual (counter)
  for row in approval_logs:
    if row["action"] == "approve":
      approved_actions = approved_actions + 1
    elif row["action"] == "reject":
      rejected_actions = rejected_actions + 1

  # ---------------------------------------------------------
  # MENGHITUNG DAN MENGELOMPOKKAN AKSI (APPROVE/REJECT) PER AKTOR
  # ---------------------------------------------------------
  actor_action_map: dict[str, dict] = {}

  for row in approval_logs:
    # 1. Menentukan Kunci Aktor (actor_key) secara eksplisit
    # Memeriksa username terlebih dahulu
    if row.get("actor_username"):
      actor_key = row.get("actor_username")
    # Jika username kosong, periksa ID karyawan
    elif row.get("actor_id_karyawan"):
      actor_key = row.get("actor_id_karyawan")
    # Jika keduanya kosong, gunakan "unknown"
    else:
      actor_key = "unknown"

    # 2. Membuat kerangka data awal jika aktor belum ada di dalam dictionary
    if actor_key not in actor_action_map:
      actor_action_map[actor_key] = {
        "actor": actor_key,
        "approve": 0,
        "reject": 0,
      }

    # 3. Menambahkan hitungan berdasarkan aksi yang dilakukan
    aksi = row["action"]
    if aksi == "approve":
      actor_action_map[actor_key]["approve"] = (
        actor_action_map[actor_key]["approve"] + 1
      )
    elif aksi == "reject":
      actor_action_map[actor_key]["reject"] = actor_action_map[actor_key]["reject"] + 1

  # ---------------------------------------------------------
  # MENGURUTKAN DATA AKTOR BERDASARKAN TOTAL AKSI
  # ---------------------------------------------------------
  def kriteria_sort_aktor(x):
    # Menghitung total aksi (jumlah approve ditambah jumlah reject)
    total_aksi = x["approve"] + x["reject"]

    # Mengembalikan dua nilai untuk diurutkan:
    # a. -total_aksi: Tanda minus berarti diurutkan dari angka paling besar (Descending)
    # b. x["actor"]: Jika total aksinya sama, urutkan berdasarkan nama abjad A-Z (Ascending)
    return (-total_aksi, x["actor"])

  # Mengambil semua nilai dari dictionary untuk diurutkan
  daftar_aktor = actor_action_map.values()
  # Melakukan pengurutan dengan fungsi bantuan yang dibuat di atas
  approval_by_actor = sorted(daftar_aktor, key=kriteria_sort_aktor)

  # SLA sederhana: selisih dari tanggal record ke timestamp aksi audit.
  absensi_lookup = {str(row["id_absensi"]): row for row in absensi_rows}
  pengajuan_lookup = {str(row["id_pengajuan"]): row for row in pengajuan_rows}
  response_hours_list = []
  for row in approval_logs:
    record_id = str(row.get("record_id") or "")
    created_at = row["created_at"]
    if row["table_name"] == "absensi":
      source_row = absensi_lookup.get(record_id)
      if not source_row:
        continue
      source_ts = source_row["tanggal_absen"]
    else:
      source_row = pengajuan_lookup.get(record_id)
      if not source_row:
        continue
      source_ts = datetime.combine(source_row["tanggal_mulai"], time.min)

    delta_hours = (created_at - source_ts).total_seconds() / 3600
    if delta_hours >= 0:
      response_hours_list.append(delta_hours)

  avg_response_hours = (
    round(
      (sum(response_hours_list) / len(response_hours_list)),
      2,
    )
    if response_hours_list
    else None
  )

  return {
    "range": {
      "start_date": start_d.isoformat(),
      "end_date": end_d.isoformat(),
      "days_count": days_count,
    },
    "overview": {
      "total_record": len(absensi_rows),
      "hadir": present_count,
      "tidak_hadir": absent_count,
      "pending": pending_count,
      "telat": late_count,
      "late_rate_percent": round((late_count / (present_count or 1)) * 100, 2),
    },
    "trend_harian": list(trend_map.values()),
    "departemen_breakdown": departemen_breakdown,
    "top_list": {
      "top_telat": top_telat,
      "top_tidak_hadir": top_tidak_hadir,
    },
    "approval_metrics": {
      "total_action": len(approval_logs),
      "approved": approved_actions,
      "rejected": rejected_actions,
      "avg_response_hours": avg_response_hours,
      "by_actor": approval_by_actor,
    },
    "audit_feed": audit_rows[:30],
  }


async def get_pengajuan(tgl: str | None = None):
  query = PengajuanAbsen.all()
  if tgl:
    query = query.filter(tanggal_mulai=date.fromisoformat(tgl))
  else:
    query = query.filter(tanggal_mulai=date.today())

  rows = await query.order_by("-id_pengajuan").values(
    "id_pengajuan",
    "id_karyawan",
    "tipe_pengajuan",
    "tanggal_mulai",
    "tanggal_akhir",
    "foto_lampiran",
    "keterangan",
    "status",
    "alasan_penolakan",
    "karyawan__nama_karyawan",
  )
  return [
    {
      "id_pengajuan": row["id_pengajuan"],
      "id_karyawan": row["id_karyawan"],
      "tipe_pengajuan": row["tipe_pengajuan"],
      "tanggal_mulai": row["tanggal_mulai"],
      "tanggal_akhir": row["tanggal_akhir"],
      "foto_lampiran": row["foto_lampiran"],
      "keterangan": row["keterangan"],
      "status": row["status"],
      "alasan_penolakan": row["alasan_penolakan"],
      "nama_karyawan": row["karyawan__nama_karyawan"],
    }
    for row in rows
  ]


def _format_str_date(params: str) -> str:
  tgl = params.split("-")
  return f"{tgl[2]}-{tgl[1]}-{tgl[0]}"


def _format_indonesian_date(date_obj: date) -> str:
  days = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
  months = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
  ]
  day_name = days[date_obj.weekday()]
  month_name = months[date_obj.month - 1]
  return f"{day_name}, {date_obj.day} {month_name} {date_obj.year}"


def _bulan_indo(month: str) -> str:
  months = {
    "01": "Januari",
    "02": "Februari",
    "03": "Maret",
    "04": "April",
    "05": "Mei",
    "06": "Juni",
    "07": "Juli",
    "08": "Agustus",
    "09": "September",
    "10": "Oktober",
    "11": "November",
    "12": "Desember",
  }
  return months.get(month, "")


async def export_excel(start_date: str | None = None, end_date: str | None = None):
  logger.info("PROSES GENERATE EXCEL REKAPITULASI")

  query = Absensi.all()
  periode_laporan = ""

  if start_date and end_date:
    start_d = datetime.strptime(start_date, "%Y-%m-%d").date()
    end_d = datetime.strptime(end_date, "%Y-%m-%d").date()
    if end_d < start_d:
      raise HTTPException(
        status_code=400, detail="Tanggal akhir tidak boleh < tanggal mulai"
      )
    periode_laporan = f"{_format_str_date(start_date)} s/d {_format_str_date(end_date)}"
    start_dt = datetime.combine(start_d, time.min)
    end_dt = datetime.combine(end_d + timedelta(days=1), time.min)
    query = query.filter(tanggal_absen__gte=start_dt, tanggal_absen__lt=end_dt)
  elif start_date:
    start_d = datetime.strptime(start_date, "%Y-%m-%d").date()
    periode_laporan = (
      _bulan_indo(start_date.split("-")[1]) + " Tahun " + str(start_d.year)
    ).upper()
    query = query.filter(tanggal_absen__month=start_d.month)
  else:
    now = datetime.now()
    periode_laporan = f"Bulan {_bulan_indo(f'{now.month:02d}')} {now.year}"
    query = query.filter(tanggal_absen__month=now.month, tanggal_absen__year=now.year)

  all_data = await query.order_by("tanggal_absen", "karyawan__nama_karyawan").values(
    "tanggal_absen",
    "karyawan__nama_karyawan",
    "karyawan__posisi",
    "check_in",
    "check_out",
    "pengajuan",
    "is_telat",
    "status_absen",
    "alasan_penolakan",
  )

  if not all_data:
    raise HTTPException(
      status_code=404,
      detail=f"Tidak ada data absensi untuk periode {periode_laporan}",
    )

  grouped_data: dict[date, list[dict]] = defaultdict(list)
  for row in all_data:
    grouped_data[row["tanggal_absen"].date()].append(row)

  wb = Workbook()
  ws = wb.active
  ws.title = "Laporan Absensi"

  ws.merge_cells("A1:I1")
  corp_cell = ws["A1"]
  corp_cell.value = "CV BENGKEL TEKNOLOGI INDONESIA"
  corp_cell.alignment = Alignment(horizontal="center", vertical="center")
  corp_cell.font = Font(bold=True, size=16)

  ws.merge_cells("A2:I2")
  ket_cell = ws["A2"]
  ket_cell.value = f"LAPORAN ABSENSI PERIODE {periode_laporan}"
  ket_cell.alignment = Alignment(horizontal="center", vertical="center")
  ket_cell.font = Font(bold=True, size=14)

  ws.append([""])

  column_headers = [
    "No",
    "Nama Karyawan",
    "Posisi",
    "Absen Masuk",
    "Absen Keluar",
    "Pengajuan",
    "Terlambat",
    "Status Absen",
    "Alasan Penolakan",
  ]

  for date_key in sorted(grouped_data.keys()):
    records_for_the_day = grouped_data[date_key]

    current_row = ws.max_row + 1
    ws.merge_cells(f"A{current_row}:I{current_row}")
    date_header_cell = ws[f"A{current_row}"]
    date_header_cell.value = _format_indonesian_date(date_key)
    date_header_cell.font = Font(bold=True, size=12)
    date_header_cell.alignment = Alignment(horizontal="left")

    ws.append(column_headers)
    header_row = ws[ws.max_row]
    for cell in header_row:
      cell.font = Font(bold=True)
      cell.alignment = Alignment(horizontal="center", vertical="center")
      cell.fill = PatternFill(start_color="D3D3D3", fill_type="solid")
      cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
      )

    for i, record in enumerate(records_for_the_day, 1):
      is_telat_str = "Ya" if record.get("is_telat") == 1 else "Tidak"
      row_data = [
        i,
        record.get("karyawan__nama_karyawan") or "-",
        record.get("karyawan__posisi") or "-",
        record["check_in"].time() if record.get("check_in") else "-",
        record["check_out"].time() if record.get("check_out") else "-",
        record.get("pengajuan") or "-",
        is_telat_str,
        record.get("status_absen") or "-",
        record.get("alasan_penolakan") or "-",
      ]
      ws.append(row_data)

    ws.append([""])

  for col_idx, column in enumerate(ws.columns, 1):
    column_letter = get_column_letter(col_idx)
    max_length = 0
    for cell in column:
      if isinstance(cell, MergedCell):
        continue
      if cell.value is None:
        continue
      max_length = max(max_length, len(str(cell.value)))

    adjusted_width = (max_length + 2) * 1.2
    if col_idx == 1:
      adjusted_width = 5
    ws.column_dimensions[column_letter].width = adjusted_width

  excel_password = "1234"
  ws.protection = SheetProtection(sheet=True)
  ws.protection.set_password(excel_password)
  ws.protection.selectLockedCells = False
  ws.protection.selectUnlockedCells = False
  ws.protection.enable()

  wb.security = WorkbookProtection(lockStructure=True)
  wb.security.set_workbook_password(excel_password)

  file_path = "data_absensi_harian.xlsx"
  if os.path.exists(file_path):
    os.chmod(file_path, 0o644)

  wb.save(file_path)
  os.chmod(file_path, 0o444)

  logger.info("SELESAI GENERATE EXCEL REKAPITULASI")

  return FileResponse(
    os.path.abspath(file_path),
    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    filename="laporan_absensi_harian.xlsx",
  )


async def update_karyawan(
  id_karyawan: str,
  payload: KaryawanUpdateRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  karyawan = await Karyawan.filter(id_karyawan=id_karyawan).first()
  if not karyawan:
    raise HTTPException(status_code=404, detail="Data karyawan tidak ditemukan")

  karyawan.nama_karyawan = payload.nama_karyawan
  karyawan.email_karyawan = payload.email_karyawan
  karyawan.nomor_hp = payload.nomor_hp
  karyawan.tanggal_rekrut = (
    date.fromisoformat(payload.tanggal_rekrut) if payload.tanggal_rekrut else None
  )
  karyawan.status = payload.status
  karyawan.id_departemen = payload.id_departemen
  karyawan.posisi = payload.posisi
  await karyawan.save()
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def update_akun(
  username: str,
  payload: AkunUpdateRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  akun = await Akun.filter(username=username).first()
  if not akun:
    raise HTTPException(status_code=404, detail="Data akun tidak ditemukan")

  akun.passwd = hashlib.md5(str(payload.passwd).encode()).hexdigest()
  akun.roles = payload.roles
  akun.id_karyawan = payload.id_karyawan
  akun.status = "aktif" if payload.status else "nonaktif"
  await akun.save()
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def update_konfigurasi(
  payload: KonfigurasiUpdateRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  konfigurasi = await KonfigurasiAplikasi.filter(
    id_pengaturan=payload.id_pengaturan
  ).first()
  if not konfigurasi:
    raise HTTPException(status_code=404, detail="Konfigurasi tidak ditemukan")

  konfigurasi.toleransi_terlambat = payload.toleransi_terlambat
  konfigurasi.maks_hari_cuti = payload.maks_hari_cuti
  await konfigurasi.save()
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def update_jadwal(
  id_jadwal: int,
  payload: JadwalUpdateRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  if not payload.shift_mulai and not payload.shift_selesai:
    raise HTTPException(status_code=400, detail="Tidak ada field jadwal yang diupdate")

  jadwal = await JadwalKerja.filter(id_jadwal=id_jadwal).first()
  if not jadwal:
    raise HTTPException(status_code=404, detail="Jadwal tidak ditemukan")

  update_fields: list[str] = []
  if payload.shift_mulai:
    jadwal.shift_mulai = time.fromisoformat(payload.shift_mulai)
    update_fields.append("shift_mulai")
  if payload.shift_selesai:
    jadwal.shift_selesai = time.fromisoformat(payload.shift_selesai)
    update_fields.append("shift_selesai")
  await jadwal.save(update_fields=update_fields)
  return {"status": "ok", "message": "Sukses Simpan Data"}


async def update_hari_libur(
  id_libur: int,
  payload: HariLiburUpdateRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  hari_libur = await HariLibur.filter(id_libur=id_libur).first()
  if not hari_libur:
    raise HTTPException(status_code=404, detail="Data hari libur tidak ditemukan")

  hari_libur.tanggal = date.fromisoformat(payload.tanggal)
  hari_libur.keterangan = payload.keterangan
  hari_libur.tipe = payload.tipe
  await hari_libur.save()
  return {"status": "ok", "message": "Sukses Update Data"}


async def update_status_absensi(
  payload: UpdateStatusAbsensiRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
  is_bulk: bool = False,
):
  if is_bulk and payload.updated_bulk_data:
    for item in payload.updated_bulk_data:
      absensi = await Absensi.filter(
        id_absensi=item["id_absensi"],
        karyawan_id=item["id_karyawan"],
      ).first()
      if not absensi:
        continue

      action_name = (
        "approve"
        if item["status_absen"] == "approved"
        else "reject"
        if item["status_absen"] == "rejected"
        else "edit"
      )
      absensi.status_absen = item["status_absen"]
      absensi.alasan_penolakan = item.get("alasan_penolakan", None)
      with audit_action(action_name, {"mode": "bulk"}):
        await absensi.save(update_fields=["status_absen", "alasan_penolakan"])

      log_message = (
        f"ADMIN  MENGUPDATE STATUS ABSENSI untuk Karyawan [{item['id_karyawan']}] "
        f"menjadi [{item['status_absen']}]"
      )
      logger.info(log_message)

      for ws_con in admin_to_user_conn:
        await ws_con.send_text(
          json.dumps(
            {
              "id_karyawan": item["id_karyawan"],
              "status": item["status_absen"],
              "message": f"Absen Anda di{item['status_absen']}",
            }
          )
        )
  else:
    absensi = await Absensi.filter(
      id_absensi=payload.id_absensi,
      karyawan_id=payload.id_karyawan,
    ).first()
    if not absensi:
      raise HTTPException(status_code=404, detail="Data absensi tidak ditemukan")

    action_name = (
      "approve"
      if payload.status_absen == "approved"
      else "reject"
      if payload.status_absen == "rejected"
      else "edit"
    )
    absensi.status_absen = payload.status_absen
    absensi.alasan_penolakan = (
      payload.alasan_penolakan if payload.alasan_penolakan else None
    )
    with audit_action(action_name, {"mode": "single"}):
      await absensi.save(update_fields=["status_absen", "alasan_penolakan"])

    log_message = (
      f"ADMIN  MENGUPDATE STATUS ABSENSI untuk Karyawan [{payload.id_karyawan}] "
      f"menjadi [{payload.status_absen}]"
    )
    logger.info(log_message)
    for ws_con in admin_to_user_conn:
      await ws_con.send_text(
        json.dumps(
          {
            "id_karyawan": payload.id_karyawan,
            "status": payload.status_absen,
            "message": f"Absen Anda di{payload.status_absen}",
          }
        )
      )

  return {"status": "ok", "message": "Sukses Update Status Absensi"}


def _daterange_inclusive(d1: date, d2: date):
  cur = d1
  while cur <= d2:
    yield cur
    cur += timedelta(days=1)


async def update_pengajuan(
  payload: UpdatePengajuanRequest,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  """
  Update_pengajuan masih punya 1 log_audit manual di admin_controller.py
  karena flow transaksinya lebih kompleks (multi operasi).
  """
  before_pengajuan = (
    await PengajuanAbsen.filter(
      id_pengajuan=payload.id_pengajuan,
      karyawan_id=payload.id_karyawan,
    )
    .limit(1)
    .values()
  )
  before_pengajuan = before_pengajuan[0] if before_pengajuan else None
  if not before_pengajuan:
    raise HTTPException(status_code=404, detail="Data pengajuan tidak ditemukan")

  async with in_transaction() as db:
    if payload.alasan_penolakan is not None:
      updated = (
        await PengajuanAbsen.filter(
          id_pengajuan=payload.id_pengajuan,
          karyawan_id=payload.id_karyawan,
        )
        .using_db(db)
        .update(
          status=payload.status,
          alasan_penolakan=payload.alasan_penolakan,
        )
      )
      if updated == 0:
        raise HTTPException(status_code=404, detail="Data pengajuan tidak ditemukan")

    elif payload.status == "approved":
      start_d = datetime.strptime(payload.tanggal_mulai, "%Y-%m-%d").date()
      end_d = datetime.strptime(payload.tanggal_akhir, "%Y-%m-%d").date()

      if end_d < start_d:
        raise HTTPException(status_code=400, detail="Tanggal Akhir < Tanggal Mulai")

      days_in_range = list(_daterange_inclusive(start_d, end_d))
      day_start = datetime.combine(start_d, time.min)
      day_end = datetime.combine(end_d + timedelta(days=1), time.min)

      existing_rows = (
        await Absensi.filter(
          karyawan_id=payload.id_karyawan,
          tanggal_absen__gte=day_start,
          tanggal_absen__lt=day_end,
        )
        .using_db(db)
        .values("id_absensi", "tanggal_absen")
      )
      existing_dates = {
        row["tanggal_absen"].date(): row["id_absensi"] for row in existing_rows
      }

      for current_date in days_in_range:
        absensi_id = existing_dates.get(current_date)
        if absensi_id:
          await (
            Absensi.filter(id_absensi=absensi_id)
            .using_db(db)
            .update(status_absen=payload.status)
          )
          continue

        check_in_date = datetime.combine(current_date, time.min)
        await Absensi.create(
          karyawan_id=payload.id_karyawan,
          tanggal_absen=check_in_date,
          check_in=check_in_date,
          latitude_checkin=0.0,
          longitude_checkin=0.0,
          foto_checkin="no-foto",
          pengajuan=payload.tipe_pengajuan,
          status_absen=payload.status,
          using_db=db,
        )

      updated = (
        await PengajuanAbsen.filter(
          id_pengajuan=payload.id_pengajuan,
          karyawan_id=payload.id_karyawan,
        )
        .using_db(db)
        .update(status=payload.status)
      )
      if updated == 0:
        raise HTTPException(status_code=404, detail="Data pengajuan tidak ditemukan")

    else:
      updated = (
        await PengajuanAbsen.filter(
          id_pengajuan=payload.id_pengajuan,
          karyawan_id=payload.id_karyawan,
        )
        .using_db(db)
        .update(status=payload.status)
      )
      if updated == 0:
        raise HTTPException(status_code=404, detail="Data pengajuan tidak ditemukan")

  after_pengajuan = (
    await PengajuanAbsen.filter(
      id_pengajuan=payload.id_pengajuan,
      karyawan_id=payload.id_karyawan,
    )
    .limit(1)
    .values()
  )
  after_pengajuan = after_pengajuan[0] if after_pengajuan else None
  await log_audit(
    actor=actor,
    action=(
      "approve"
      if payload.status == "approved"
      else "reject"
      if payload.status == "rejected"
      else "edit"
    ),
    table_name="pengajuan_absen",
    record_id=payload.id_pengajuan,
    before_data=before_pengajuan,
    after_data=after_pengajuan,
  )

  log_message = (
    f"ADMIN MENGUPDATE STATUS PENGAJUAN untuk Karyawan [{payload.id_karyawan}] "
    f"menjadi [{payload.status}]"
  )
  logger.info(log_message)

  for ws_con in admin_to_user_conn:
    await ws_con.send_text(
      json.dumps(
        {
          "id_karyawan": payload.id_karyawan,
          "status": payload.status,
          "message": f"Pengajuan Anda telah di-{payload.status}",
        }
      )
    )

  return {"status": "ok", "message": "Sukses Update Pengajuan"}


async def unbind_device(
  username: str,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  akun = await Akun.filter(username=username).first()
  if not akun:
    raise HTTPException(status_code=404, detail="Akun tidak ditemukan")

  akun.device_id = None
  with audit_action("edit", {"activity": "unbind_device"}):
    await akun.save(update_fields=["device_id"])
  return {"status": "ok", "message": "Sukses Unbind Device Data"}


async def delete_karyawan(
  id_karyawan: str,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  karyawan = await Karyawan.filter(id_karyawan=id_karyawan).first()
  if not karyawan:
    raise HTTPException(status_code=404, detail="Data karyawan tidak ditemukan")
  await karyawan.delete()
  return {"status": "ok", "message": "Sukses Delete Data"}


async def delete_akun(
  username: str,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  akun = await Akun.filter(username=username).first()
  if not akun:
    raise HTTPException(status_code=404, detail="Data akun tidak ditemukan")
  await akun.delete()
  return {"status": "ok", "message": "Sukses Delete Data"}


async def delete_departemen(
  id_departemen: int,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  departemen = await Departemen.filter(id_departemen=id_departemen).first()
  if not departemen:
    raise HTTPException(status_code=404, detail="Data departemen tidak ditemukan")
  await departemen.delete()
  return {"status": "ok", "message": "Sukses Delete Data"}


async def delete_hari_libur(
  id_libur: int,
  actor: JwtAuthorizationCredentials | dict | None = None,
) -> dict:
  hari_libur = await HariLibur.filter(id_libur=id_libur).first()
  if not hari_libur:
    raise HTTPException(status_code=404, detail="Data hari libur tidak ditemukan")
  await hari_libur.delete()
  return {"status": "ok", "message": "Sukses Delete Data"}
